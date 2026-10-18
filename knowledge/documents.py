"""Extract each upload independently, retaining PDF page locations."""

import hashlib
import io
from pathlib import PurePosixPath
import zipfile


def extract_upload(upload):
    filename = PurePosixPath((upload.filename or "upload").replace("\\", "/")).name[:250]
    extension = filename.rsplit(".", 1)[-1].lower()
    data = upload.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        raise ValueError("Files must be smaller than 20 MB.")
    metadata = {"filename": filename}
    try:
        if extension == "pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
            parts, spans, offset = [], [], 0
            for index, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if parts:
                    offset += 2
                parts.append(text)
                spans.append({"start": offset, "end": offset + len(text), "page": index})
                offset += len(text)
            text = "\n\n".join(parts)
            metadata["page_spans"] = spans
        elif extension == "docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("The expanded document is too large.")
            document = Document(io.BytesIO(data))
            # Include table cells as well as paragraphs; both may contain source evidence.
            text = "\n\n".join([paragraph.text for paragraph in document.paragraphs]
                                  + [" | ".join(cell.text for cell in row.cells)
                                     for table in document.tables for row in table.rows])
        elif extension == "doc":
            raise ValueError("Please convert legacy .doc files to .docx before uploading.")
        elif extension in {"txt", "md", "json", "csv"}:
            text = data.decode("utf-8-sig")
        else:
            raise ValueError("Supported files are PDF, DOCX, TXT, Markdown, JSON, and CSV.")
    except ValueError:
        raise
    except Exception:
        raise ValueError(f"Could not extract text from {filename}. Check the file format.") from None
    if not text.strip():
        raise ValueError(f"No text was extracted from {filename}. Scanned PDFs need OCR before upload.")
    identity = filename.encode() + b"\0" + data
    return {"document_id": "file_" + hashlib.sha256(identity).hexdigest()[:24],
            "text": text, "title": filename, "source_type": extension, "metadata": metadata}
